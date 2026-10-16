from datetime import datetime, date
from decimal import Decimal, InvalidOperation
from django.core.management.base import BaseCommand, CommandError
from openpyxl import load_workbook
from wip.models import OrdenProducto, MovimientoWIP

class Command(BaseCommand):
    help = 'Importa movimientos de DB_HISTORIAL desde el Excel original. Ejecutar primero en una base de datos vacía o revisar duplicados.'
    def add_arguments(self, parser):
        parser.add_argument('archivo', type=str, help='Ruta al archivo XLSX')
        parser.add_argument('--skip-existing', action='store_true', help='Omite movimientos que coincidan por fecha, tipo, ID, referencia y cantidades')
    def handle(self, *args, **opts):
        try: wb = load_workbook(opts['archivo'], data_only=True, read_only=True)
        except Exception as e: raise CommandError(f'No se pudo abrir el Excel: {e}')
        if 'DB_HISTORIAL' not in wb.sheetnames: raise CommandError('No se encontró la hoja DB_HISTORIAL')
        ws = wb['DB_HISTORIAL']; headers = [str(c.value).strip() if c.value is not None else '' for c in ws[1]]
        ix = {h.upper():i for i,h in enumerate(headers)}
        required = ['FECHA','TIPO MOVIMIENTO','ID','DETALLE PRODUCTO','CANTIDAD (KG)','UFE','SCRAP (KG)','AREA ORIGEN','AREA DESTINO','AREA POSTERIOR','ENTREGADO POR','RECIBIDO POR','REF/HT','OBSERVACIONES','COSTO (BS/KG)']
        missing = [h for h in required if h not in ix]
        if missing: raise CommandError('Faltan columnas: '+', '.join(missing))
        created_products = created_moves = skipped = errors = 0
        def val(row, key): return row[ix[key]] if ix[key] < len(row) else None
        def dec(x, default='0'):
            if x in (None,''): return Decimal(default)
            if isinstance(x,str): x=x.replace(',','.')
            return Decimal(str(x))
        for rownum, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if not any(x is not None for x in row): continue
            try:
                pid_raw = val(row,'ID')
                if pid_raw is None: continue
                pid = str(int(pid_raw)) if isinstance(pid_raw,(int,float)) and float(pid_raw).is_integer() else str(pid_raw).strip()
                rawdate = val(row,'FECHA')
                if isinstance(rawdate, datetime): fecha=rawdate.date()
                elif isinstance(rawdate,date): fecha=rawdate
                else: fecha=date.fromisoformat(str(rawdate)[:10])
                tipo = str(val(row,'TIPO MOVIMIENTO') or '').strip().upper()
                if tipo not in ('INGRESO','SALIDA','TRASPASO'): raise ValueError(f'Tipo no válido: {tipo}')
                detalle = str(val(row,'DETALLE PRODUCTO') or '').strip()
                prod, made = OrdenProducto.objects.get_or_create(id_producto=pid, defaults={'descripcion_referencia':detalle})
                if made: created_products += 1
                elif not prod.descripcion_referencia and detalle: prod.descripcion_referencia=detalle; prod.save(update_fields=['descripcion_referencia'])
                ref = str(val(row,'REF/HT') or '').strip()
                kg, ufe, scrap = dec(val(row,'CANTIDAD (KG)')), dec(val(row,'UFE')), dec(val(row,'SCRAP (KG)'))
                if opts['skip_existing'] and MovimientoWIP.objects.filter(fecha=fecha,tipo_movimiento=tipo,producto=prod,referencia_ht=ref,cantidad_kg=kg,ufe=ufe).exists(): skipped+=1; continue
                cost = val(row,'COSTO (BS/KG)')
                Movement = MovimientoWIP.objects.create(fecha=fecha,tipo_movimiento=tipo,producto=prod,detalle_producto=detalle,cantidad_kg=kg,ufe=ufe,scrap_kg=scrap,area_origen=str(val(row,'AREA ORIGEN') or '').strip(),area_destino=str(val(row,'AREA DESTINO') or '').strip(),area_posterior=str(val(row,'AREA POSTERIOR') or '').strip(),entregado_por=str(val(row,'ENTREGADO POR') or '').strip(),recibido_por=str(val(row,'RECIBIDO POR') or '').strip(),referencia_ht=ref,observaciones=str(val(row,'OBSERVACIONES') or '').strip(),costo_bs_kg=dec(cost) if cost not in (None,'') else None)
                created_moves += 1
            except Exception as e:
                errors += 1; self.stderr.write(f'Fila {rownum}: {e}')
        self.stdout.write(self.style.SUCCESS(f'Productos creados: {created_products}; movimientos importados: {created_moves}; omitidos: {skipped}; errores: {errors}'))
